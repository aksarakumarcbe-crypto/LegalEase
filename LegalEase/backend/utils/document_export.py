from __future__ import annotations
from io import BytesIO
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF
from backend.utils.text_utils import sanitize_text

def _split_lines(text):
    return [line.rstrip() for line in sanitize_text(text).splitlines()]

def _is_heading(line):
    s = line.strip()
    return bool(s) and len(s) < 100 and (
        s.isupper() or s.startswith(tuple(f"{n}." for n in range(1, 100))) or s.endswith(":")
    )

def format_docx(text, doc_type="Legal Document", company_name="LegalEase"):
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(.7)
    section.bottom_margin = Inches(.7)
    section.left_margin = Inches(.8)
    section.right_margin = Inches(.8)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    p = section.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(company_name or "LegalEase")
    r.bold = True
    r.font.size = Pt(14)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run(doc_type.upper())
    r.bold = True
    r.font.size = Pt(16)

    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("AI-GENERATED FIRST DRAFT")
    r.italic = True
    r.font.size = Pt(9)

    for line in _split_lines(text):
        s = line.strip()
        if not s:
            document.add_paragraph()
            continue
        if s.startswith(("• ", "- ", "* ")):
            p = document.add_paragraph(style="List Bullet")
            p.add_run(s.lstrip("•-* ").strip())
            continue
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(s)
        if _is_heading(s):
            r.bold = True
            r.font.size = Pt(12)

    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(
        "LegalEase — AI-generated draft. Review with a qualified legal professional."
    )
    r.font.size = Pt(8)

    out = BytesIO()
    document.save(out)
    return out.getvalue()

class LegalPDF(FPDF):
    def __init__(self, company_name):
        super().__init__()
        self.company_name = company_name or "LegalEase"
        self.set_auto_page_break(auto=True, margin=18)

    def header(self):
        self.set_font("Times", "B", 13)
        self.cell(0, 8, self.company_name, new_x="LMARGIN", new_y="NEXT", align="C")
        self.line(15, 18, 195, 18)
        self.ln(4)

    def footer(self):
        self.set_y(-14)
        self.set_font("Times", "", 8)
        self.cell(
            0, 8,
            f"LegalEase | Page {self.page_no()} | AI-generated draft - review before use.",
            align="C"
        )

def format_pdf(text, doc_type="Legal Document", company_name="LegalEase"):
    pdf = LegalPDF(company_name)
    pdf.set_margins(15, 24, 15)
    pdf.add_page()
    pdf.set_font("Times", "B", 16)
    pdf.multi_cell(0, 9, doc_type.upper(), align="C", new_x="LMARGIN")
    pdf.set_font("Times", "I", 9)
    pdf.multi_cell(0, 6, "AI-GENERATED FIRST DRAFT", align="C", new_x="LMARGIN")
    pdf.ln(4)

    for line in _split_lines(text):
        s = line.strip()
        if not s:
            pdf.ln(3)
            continue
        if s.startswith(("• ", "- ", "* ")):
            pdf.set_font("Times", "", 11)
            pdf.multi_cell(0, 6, "- " + s.lstrip("•-* ").strip(), new_x="LMARGIN")
            continue
        pdf.set_font("Times", "B" if _is_heading(s) else "", 11)
        pdf.multi_cell(0, 6, s, new_x="LMARGIN")

    return bytes(pdf.output(dest="S"))

def format_txt(text):
    return sanitize_text(text).encode("utf-8")
